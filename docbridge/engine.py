"""Bounded conversion operations. Called only inside a disposable worker process."""
import io
import os
import subprocess
import warnings
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageOps
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from pypdf import PdfReader
import pypdfium2 as pdfium
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PDFImage

MAX_PAGES = 50
MAX_PIXELS = 25_000_000
MAX_TEXT = 1_000_000
MAX_OUTPUT = 40 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = MAX_PIXELS
warnings.simplefilter('error', Image.DecompressionBombWarning)
FORMATS = {
    'txt': ['pdf', 'docx'],
    'docx': ['pdf', 'txt'],
    'pdf': ['txt', 'docx', 'png'],
    'jpg': ['pdf', 'txt', 'docx'],
    'jpeg': ['pdf', 'txt', 'docx'],
    'png': ['pdf', 'txt', 'docx'],
}


class ConversionError(Exception):
    pass


def command(args, timeout=60):
    try:
        result = subprocess.run(args, stdout=subprocess.DEVNULL,
                                stderr=subprocess.DEVNULL, timeout=timeout, check=False)
    except FileNotFoundError:
        raise ConversionError('A required conversion engine is unavailable.') from None
    except subprocess.TimeoutExpired:
        raise ConversionError('Conversion engine timed out.') from None
    if result.returncode:
        raise ConversionError('The conversion engine could not process this document.')


def clean_text(text):
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    # XML 1.0 compatible text for both DOCX and ReportLab paragraphs.
    text = ''.join(c for c in text if c in '\n\t' or
                   0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or
                   0x10000 <= ord(c) <= 0x10FFFF)
    if len(text) > MAX_TEXT:
        raise ConversionError('Extracted text exceeds the one-million-character limit.')
    return text


def read_utf8(path):
    try:
        return clean_text(path.read_text(encoding='utf-8-sig'))
    except UnicodeError:
        raise ConversionError('TXT input must use UTF-8 encoding.') from None


def read_image(path, source):
    try:
        with Image.open(path) as image:
            expected = 'JPEG' if source in ('jpg', 'jpeg') else 'PNG'
            if image.format != expected or getattr(image, 'n_frames', 1) != 1:
                raise ConversionError('File content does not match its format, or image is animated.')
            if image.width * image.height > MAX_PIXELS:
                raise ConversionError('Image exceeds 25 megapixels.')
            image.load()
            oriented = ImageOps.exif_transpose(image).convert('RGBA')
            white = Image.new('RGBA', oriented.size, 'white')
            white.alpha_composite(oriented)
            return white.convert('RGB')
    except (Image.DecompressionBombError, Image.DecompressionBombWarning):
        raise ConversionError('Image exceeds the pixel limit.') from None


def validate_docx(path):
    """Reject oversized archives, active content and relationships that fetch remote data."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
            raise ConversionError('DOCX archive exceeds the unpacked size limit.')
        names = [e.filename for e in entries]
        if len(set(names)) != len(names) or 'word/document.xml' not in names:
            raise ConversionError('Invalid DOCX archive.')
        for entry in entries:
            name = entry.filename
            lower = name.lower()
            if (name.startswith('/') or '\\' in name or '..' in name.split('/') or
                entry.flag_bits & 1 or any(x in lower for x in ('vbaproject', 'embeddings/', 'activex/'))):
                raise ConversionError('DOCX contains unsupported embedded or active content.')
            if entry.file_size > 1024 * 1024 and entry.file_size > 250 * max(entry.compress_size, 1):
                raise ConversionError('DOCX compression ratio exceeds the limit.')
            if lower.endswith(('.xml', '.rels')):
                parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
                root = etree.fromstring(archive.read(name), parser)
                if root.getroottree().docinfo.doctype:
                    raise ConversionError('XML document types are not supported.')
                for node in root.iter():
                    if node.get('TargetMode', '').lower() == 'external':
                        raise ConversionError('Remove external links or linked resources from the DOCX first.')
                    if etree.QName(node).localname in ('instrText', 'fldSimple'):
                        instruction = (node.text or '') + ' ' + (node.get(qn('w:instr')) or '')
                        if any(x in instruction.upper() for x in ('DDE', 'INCLUDETEXT', 'INCLUDEPICTURE', 'LINK ')):
                            raise ConversionError('DOCX contains an unsupported external field.')


def docx_text(path):
    document = Document(path)
    lines = []
    for node in document.element.body:
        if node.tag == qn('w:p'):
            lines.append(''.join(t.text or '' for t in node.iter(qn('w:t'))))
        elif node.tag == qn('w:tbl'):
            for row in node.iter(qn('w:tr')):
                lines.append('\t'.join(' '.join(t.text or '' for t in cell.iter(qn('w:t')))
                                       for cell in row.findall(qn('w:tc'))))
    return clean_text('\n'.join(lines))


def font_path():
    candidates = [os.getenv('DOCBRIDGE_FONT', ''),
                  '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                  'C:/Windows/Fonts/arial.ttf']
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise ConversionError('A Unicode font is required to create PDFs.')


def text_pdf(text, output):
    pdfmetrics.registerFont(TTFont('DocBridge', font_path()))
    style = ParagraphStyle('Body', fontName='DocBridge', fontSize=10, leading=15,
                           spaceAfter=5, splitLongWords=True)
    story = []
    for line in clean_text(text).split('\n'):
        # Bound a single paragraph so exceptionally long lines can split across pages.
        if not line:
            story.append(Spacer(1, 8))
        for offset in range(0, len(line), 4000):
            story.append(Paragraph(escape(line[offset:offset + 4000]).replace('\t', '    '), style))
    if not story:
        story = [Spacer(1, 8)]
    doc = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=42, leftMargin=42,
                            topMargin=42, bottomMargin=42, title='Converted document')
    doc.build(story)
    if len(PdfReader(output).pages) > MAX_PAGES:
        raise ConversionError('The generated document exceeds 50 pages.')


def text_docx(text, output):
    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'DejaVu Sans'
    normal.font.size = Pt(11)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.left_margin = section.right_margin = Inches(0.75)
    for line in clean_text(text).split('\n'):
        document.add_paragraph(line)
    document.save(output)


def ocr_image(image, work, language):
    raster = work / 'ocr-input.png'
    image.save(raster)
    output_base = work / 'ocr-result'
    command(['tesseract', str(raster), str(output_base), '-l', language, '--psm', '3', 'txt'])
    return read_utf8(output_base.with_suffix('.txt'))


def pdf_pages(path):
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ConversionError('Password-protected PDFs are not supported.')
    if not 1 <= len(reader.pages) <= MAX_PAGES:
        raise ConversionError('PDF must contain between 1 and 50 pages.')
    return reader


def raster_page(doc, index):
    page = doc[index]
    try:
        width, height = page.get_size()
        if not 0 < width <= 14400 or not 0 < height <= 14400:
            raise ConversionError('PDF page dimensions are unsupported.')
        scale = min(200 / 72, (MAX_PIXELS / (width * height)) ** 0.5)
        bitmap = page.render(scale=scale)
        try:
            return bitmap.to_pil().convert('RGB').copy()
        finally:
            bitmap.close()
    finally:
        page.close()


def pdf_text(path, work, language, mode):
    reader = pdf_pages(path)
    lines, count, ocr_pages = [], 0, 0
    with pdfium.PdfDocument(str(path)) as raster:
        for i, page in enumerate(reader.pages):
            text = '' if mode == 'always' else (page.extract_text() or '')
            if mode == 'always' or (mode == 'auto' and not text.strip()):
                image = raster_page(raster, i)
                try:
                    text = ocr_image(image, work, language)
                finally:
                    image.close()
                ocr_pages += 1
            count += len(text)
            if count > MAX_TEXT:
                raise ConversionError('Extracted text exceeds the limit.')
            lines.append(text)
    return clean_text('\n\n'.join(lines)), ocr_pages


def convert(path, source, target, language, mode, work):
    if target not in FORMATS.get(source, []):
        raise ConversionError('Unsupported conversion pair.')
    if language not in ('eng', 'rus', 'eng+rus') or mode not in ('auto', 'always', 'never'):
        raise ConversionError('Invalid OCR language or mode.')
    output = work / ('result.zip' if target == 'png' else 'result.' + target)
    metadata = {'ocr_pages': 0, 'warnings': []}
    if source == 'docx':
        validate_docx(path)
        if target == 'pdf':
            profile = (work / 'lo-profile').resolve().as_uri()
            command(['soffice', '-env:UserInstallation=' + profile, '--headless', '--norestore',
                     '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(work), str(path)], timeout=90)
            generated = work / (path.stem + '.pdf')
            if not generated.is_file():
                raise ConversionError('LibreOffice did not produce a PDF.')
            generated.replace(output)
            pdf_pages(output)
        else:
            output.write_text(docx_text(path), encoding='utf-8')
            metadata['warnings'].append('Body paragraphs and tables only; headers, footnotes and images are omitted.')
    elif source in ('jpg', 'jpeg', 'png'):
        image = read_image(path, source)
        try:
            if target == 'pdf':
                raster = work / 'normalized.png'
                image.save(raster)
                width, height = image.size
                scale = min((A4[0] - 72) / width, (A4[1] - 72) / height)
                SimpleDocTemplate(str(output), pagesize=A4, topMargin=30, bottomMargin=30,
                                  leftMargin=30, rightMargin=30).build([
                                      PDFImage(str(raster), width=width * scale, height=height * scale)])
            else:
                if mode == 'never':
                    raise ConversionError('Image-to-text conversion requires OCR.')
                text = ocr_image(image, work, language)
                metadata['ocr_pages'] = 1
                if target == 'txt':
                    output.write_text(text, encoding='utf-8')
                else:
                    text_docx(text, output)
        finally:
            image.close()
    elif source == 'pdf' and target == 'png':
        pdf_pages(path)
        with pdfium.PdfDocument(str(path)) as raster, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
            total = 0
            for i in range(len(raster)):
                image = raster_page(raster, i)
                try:
                    buffer = io.BytesIO()
                    image.save(buffer, format='PNG')
                    total += buffer.tell()
                    if total > MAX_OUTPUT:
                        raise ConversionError('Rendered images exceed the 40 MiB output limit.')
                    archive.writestr('page-%03d.png' % (i + 1), buffer.getvalue())
                finally:
                    image.close()
    else:
        if source == 'pdf':
            text, metadata['ocr_pages'] = pdf_text(path, work, language, mode)
            metadata['warnings'].append('Text extraction does not preserve original page layout.')
        else:
            text = read_utf8(path)
        if target == 'pdf':
            text_pdf(text, output)
        elif target == 'docx':
            text_docx(text, output)
        else:
            output.write_text(text, encoding='utf-8')
    if not output.is_file() or output.stat().st_size > MAX_OUTPUT:
        raise ConversionError('Output exceeds the 40 MiB limit.')
    return output, metadata
