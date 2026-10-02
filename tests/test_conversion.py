import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from pypdf import PdfReader, PdfWriter
from docbridge import api, engine


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.work = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def convert(self, source, target, data, mode='auto'):
        path = self.work / ('input.' + source)
        path.write_bytes(data)
        return engine.convert(path, source, target, 'eng', mode, self.work)

    def test_utf8_pdf_roundtrip(self):
        text = 'Document conversion\nEnglish and Russian: Привет мир\n<literal> & symbols'
        out, _ = self.convert('txt', 'pdf', text.encode())
        extracted = ''.join(p.extract_text() for p in PdfReader(out).pages)
        self.assertIn('Привет мир', extracted)
        self.assertIn('<literal> & symbols', extracted)
        text_out, meta = self.convert('pdf', 'txt', out.read_bytes(), 'never')
        self.assertIn('Document conversion', text_out.read_text(encoding='utf-8'))
        self.assertEqual(meta['ocr_pages'], 0)

    def test_txt_docx(self):
        out, _ = self.convert('txt', 'docx', b'First paragraph\nSecond paragraph')
        self.assertEqual([p.text for p in Document(out).paragraphs], ['First paragraph', 'Second paragraph'])

    def test_docx_text_preserves_table_order(self):
        doc = Document()
        doc.add_paragraph('Before')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Name'
        table.cell(0, 1).text = 'Value'
        doc.add_paragraph('After')
        buf = io.BytesIO()
        doc.save(buf)
        out, _ = self.convert('docx', 'txt', buf.getvalue())
        self.assertEqual(out.read_text(), 'Before\nName\tValue\nAfter')

    def test_pdf_docx(self):
        pdf, _ = self.convert('txt', 'pdf', b'Editable content')
        out, meta = self.convert('pdf', 'docx', pdf.read_bytes(), 'never')
        self.assertIn('Editable content', '\n'.join(p.text for p in Document(out).paragraphs))
        self.assertTrue(meta['warnings'])

    def test_transparent_image_pdf(self):
        buf = io.BytesIO()
        Image.new('RGBA', (640, 480), (20, 100, 200, 128)).save(buf, format='PNG')
        out, _ = self.convert('png', 'pdf', buf.getvalue())
        self.assertEqual(len(PdfReader(out).pages), 1)

    def test_pdf_pages_zip(self):
        pdf, _ = self.convert('txt', 'pdf', b'Page image')
        out, _ = self.convert('pdf', 'png', pdf.read_bytes())
        with zipfile.ZipFile(out) as archive:
            self.assertEqual(archive.namelist(), ['page-001.png'])
            with Image.open(io.BytesIO(archive.read('page-001.png'))) as image:
                self.assertGreater(image.width, 1000)

    def test_reject_invalid_utf8(self):
        with self.assertRaisesRegex(engine.ConversionError, 'UTF-8'):
            self.convert('txt', 'pdf', b'\xff\xfe\x00')

    def test_reject_encrypted_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt('secret')
        buf = io.BytesIO()
        writer.write(buf)
        with self.assertRaisesRegex(engine.ConversionError, 'Password'):
            self.convert('pdf', 'txt', buf.getvalue())

    def test_reject_page_limit(self):
        writer = PdfWriter()
        for _ in range(51):
            writer.add_blank_page(width=100, height=100)
        buf = io.BytesIO()
        writer.write(buf)
        with self.assertRaisesRegex(engine.ConversionError, '50 pages'):
            self.convert('pdf', 'txt', buf.getvalue())

    def test_reject_mismatched_image(self):
        buf = io.BytesIO()
        Image.new('RGB', (10, 10)).save(buf, format='PNG')
        with self.assertRaises(engine.ConversionError):
            self.convert('jpg', 'pdf', buf.getvalue())

    def test_reject_external_docx(self):
        doc = Document()
        doc.part.relate_to('https://example.com/private',
                          'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                          is_external=True)
        buf = io.BytesIO()
        doc.save(buf)
        with self.assertRaisesRegex(engine.ConversionError, 'external'):
            self.convert('docx', 'txt', buf.getvalue())

    def require_engine(self, name):
        if not shutil.which(name):
            if os.getenv('REQUIRE_ENGINES') == '1':
                self.fail(name + ' is required in the Docker image')
            self.skipTest(name + ' is not installed locally')

    def test_libreoffice_pdf(self):
        self.require_engine('soffice')
        doc = Document()
        doc.add_paragraph('Office conversion test')
        doc.save(self.work / 'input.docx')
        out, _ = engine.convert(self.work / 'input.docx', 'docx', 'pdf', 'eng', 'auto', self.work)
        self.assertIn('Office conversion test', PdfReader(out).pages[0].extract_text())

    def test_real_ocr_and_scanned_pdf(self):
        self.require_engine('tesseract')
        image = Image.new('RGB', (1600, 450), 'white')
        draw = ImageDraw.Draw(image)
        draw.text((65, 100), 'DOCUMENT TEST 12345', font=ImageFont.truetype(engine.font_path(), 65), fill='black')
        buf = io.BytesIO()
        image.save(buf, format='PNG')
        out, meta = self.convert('png', 'txt', buf.getvalue())
        self.assertIn('DOCUMENT TEST 12345', out.read_text())
        self.assertEqual(meta['ocr_pages'], 1)
        pdf, _ = self.convert('png', 'pdf', buf.getvalue())
        out, meta = self.convert('pdf', 'txt', pdf.read_bytes())
        self.assertIn('DOCUMENT TEST 12345', out.read_text())
        self.assertEqual(meta['ocr_pages'], 1)


class APITests(unittest.TestCase):
    def call(self, *, body=b'Hello API', query='source=txt&target=pdf', key='x' * 40,
             length=None, method='POST', content_type='application/octet-stream', path='/v1/convert'):
        captured = {}
        def start(status, headers):
            captured.update(status=int(status.split()[0]), headers=dict(headers))
        env = {'PATH_INFO': path, 'REQUEST_METHOD': method, 'QUERY_STRING': query,
               'HTTP_AUTHORIZATION': 'Bearer ' + key, 'CONTENT_TYPE': content_type,
               'CONTENT_LENGTH': str(len(body)) if length is None else length,
               'wsgi.input': io.BytesIO(body)}
        with patch.dict(os.environ, {'DOCBRIDGE_API_KEY': 'x' * 40}):
            captured['body'] = b''.join(api.application(env, start))
        return captured

    def test_auth_before_conversion(self):
        with patch.object(api, 'run_job') as job:
            self.assertEqual(self.call(key='wrong')['status'], 401)
            job.assert_not_called()

    def test_upload_limit(self):
        self.assertEqual(self.call(length=str(api.MAX_UPLOAD + 1))['status'], 413)

    def test_missing_length(self):
        self.assertEqual(self.call(length='')['status'], 411)

    def test_truncated_upload(self):
        self.assertEqual(self.call(length='100')['status'], 400)

    def test_bad_options(self):
        for query in ('source=exe&target=pdf', 'source=txt&target=pdf&target=docx',
                      'source=pdf&target=txt&language=../../../secret'):
            self.assertEqual(self.call(query=query)['status'], 400)

    def test_raw_upload_required(self):
        self.assertEqual(self.call(content_type='multipart/form-data')['status'], 415)

    def test_formats(self):
        response = self.call(path='/v1/formats', method='GET')
        self.assertEqual(response['status'], 200)
        self.assertEqual(json.loads(response['body'])['formats'], engine.FORMATS)

    def test_real_worker_and_cleanup(self):
        visited = []
        original = api.run_job
        def track(work):
            visited.append(work)
            return original(work)
        with patch.object(api, 'run_job', track):
            response = self.call()
        self.assertEqual(response['status'], 200, response['body'])
        self.assertTrue(response['body'].startswith(b'%PDF'))
        self.assertEqual(response['headers']['Cache-Control'], 'no-store')
        self.assertTrue(visited)
        self.assertFalse(visited[0].exists())

    def test_worker_failure_cleanup(self):
        visited = []
        def fail(work):
            visited.append(work)
            return {'ok': False, 'error': 'Invalid document.'}
        with patch.object(api, 'run_job', fail):
            self.assertEqual(self.call()['status'], 422)
        self.assertFalse(visited[0].exists())


if __name__ == '__main__':
    unittest.main()
